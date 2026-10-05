"""Text extraction + local LLM summary (Ollama). No cloud, no API keys."""

import hashlib
import io
import json
import os
import re
import httpx
import ollama
import pdfplumber
from docx import Document as DocxDocument
from dotenv import load_dotenv
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from src.search.skills import merge_skills
from src.ai import models
from src.services import timing  # no project imports inside it, so no import cycle

load_dotenv()

OLLAMA_MODEL = os.getenv("OLLAMA_MODEL", "phi4-mini")
OLLAMA_HOST = os.getenv("OLLAMA_HOST", "http://localhost:11434")

OLLAMA_NUM_CTX = int(os.getenv("OLLAMA_NUM_CTX", "6144"))
OLLAMA_KEEP_ALIVE = os.getenv("OLLAMA_KEEP_ALIVE", "30m")
OLLAMA_MAX_OUTPUT = int(os.getenv("OLLAMA_MAX_OUTPUT", "600"))
OLLAMA_TIMEOUT = float(os.getenv("OLLAMA_TIMEOUT", "180"))
OLLAMA_NUM_THREAD = int(os.getenv("OLLAMA_NUM_THREAD", "0"))  # 0 = let Ollama decide

# Only a guard against huge files: normal resumes and one-sliders are read completely.
MAX_SLIDER_CHARS = int(os.getenv("MAX_SLIDER_CHARS", "6000"))
MAX_RESUME_CHARS = int(os.getenv("MAX_RESUME_CHARS", "9000"))

# Change when the prompts change: cached results (input hashes) are then redone.
PROMPT_VERSION = "v3"

# Per-call output limits (a tighter limit also stops runaway generation).
SUMMARY_MAX_OUTPUT = 400
PROJECTS_MAX_OUTPUT = 400
MAX_CERTS_READ = 10

# OLLAMA_NUM_CTX = int(os.getenv("OLLAMA_NUM_CTX", "8192"))
# OLLAMA_TIMEOUT = float(os.getenv("OLLAMA_TIMEOUT", "180"))

# # Keep the prompt inside the context window (roughly 4 chars per token).
# MAX_FEEDBACK_CHARS = 5000
# MAX_SLIDER_CHARS = 5000
# MAX_RESUME_CHARS = 6000


class SummaryError(Exception):
    """Raised when a summary cannot be produced (message is user-friendly)."""


# --------------------------------------------------------------------------
# Text extraction
# --------------------------------------------------------------------------
def extract_text_from_pdf(file_bytes: bytes) -> str:
    pages = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text() or "")
    return "\n".join(pages).strip()


def extract_text_from_docx(file_bytes: bytes) -> str:
    doc = DocxDocument(io.BytesIO(file_bytes))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _shape_texts(shape):
    """Yield text from a shape, recursing into groups and reading tables."""
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            yield from _shape_texts(child)
        return
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if text:
            yield text
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                yield " | ".join(cells)


def extract_text_from_pptx(file_bytes: bytes) -> str:
    presentation = Presentation(io.BytesIO(file_bytes))
    parts = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            parts.extend(_shape_texts(shape))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Notes: {notes}")
    return "\n".join(parts).strip()


def extract_text(file_name: str, file_bytes: bytes) -> str:
    ext = os.path.splitext(file_name)[1].lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_bytes)
    if ext == ".docx":
        return extract_text_from_docx(file_bytes)
    if ext == ".pptx":
        return extract_text_from_pptx(file_bytes)
    if ext in (".ppt", ".doc"):
        raise ValueError(f"Old '{ext}' files are not supported. Save as .pptx/.docx.")
    raise ValueError(f"Unsupported file type '{ext}'. Use PDF, DOCX or PPTX.")

# --------------------------------------------------------------------------
# Text preparation
# --------------------------------------------------------------------------
_BULLET = re.compile(r"^[\s•●▪◦‣·\-–—*]+(?=\S)")


def clean_text(text) -> str:
    """Collapse spaces, drop empty or symbol-only lines and leading bullet marks (fewer tokens)."""
    lines = []
    for raw in (text or "").splitlines():
        line = _BULLET.sub("", " ".join(raw.split()))
        if line and re.search(r"\w", line):
            lines.append(line)
    return "\n".join(lines)


def _prepare(text, limit: int) -> str:
    return clean_text(text)[:limit]


def input_hash(resume_text, one_slider_text, model=None) -> str:
    """Fingerprint of everything the AI reads. Same fingerprint = the AI result is still valid."""
    model = model or models.get_summary_model()
    blob = "\n".join(
        [PROMPT_VERSION, model, _prepare(one_slider_text, MAX_SLIDER_CHARS), _prepare(resume_text, MAX_RESUME_CHARS)]
    )
    return hashlib.sha256(blob.encode("utf-8")).hexdigest()


def _words(text: str) -> list[str]:
    return re.findall(r"[a-z0-9+#.]+", (text or "").lower())


# --------------------------------------------------------------------------
# Prompts. Both calls share the same system message and the same documents block,
# and differ only in the TASK at the end, so Ollama can reuse the already-read text.
# --------------------------------------------------------------------------
SYSTEM_PROMPT = (
    "You are an HR assistant writing profiles of bench employees for project managers. "
    "Use ONLY the information in the documents. Never invent skills, employers, projects, "
    "titles, dates or facts. If something is missing, leave it empty. "
    "Priority when the documents overlap or disagree: 1) ORGANIZATION ONE-SLIDER, 2) RESUME. "
    "Respond with JSON only, containing exactly the fields requested in the TASK."
)

DOCUMENTS_TEMPLATE = """### ORGANIZATION ONE-SLIDER (first priority)
{one_slider}

### RESUME (second priority)
{resume}

"""

SUMMARY_TASK = """### TASK
Return JSON with these fields:
1. headline: the professional title shown directly under the person's name at the top of the resume, or the title on the one-slider (for example "Data Scientist"). Empty string if there is none.
2. current_role: the person's most recent job title in the experience section (for example "Associate Manager"). Empty string if there is none.
3. organization_summary: 3 sentences on what this person has done INSIDE OUR ORGANIZATION (roles, projects, contributions). Use the one-slider first and the resume for background. If there is no one-slider, say so and summarise the resume.
4. skills: up to 20 short skill names (tools, technologies, domains), 1-4 words each, from both documents. No sentences, no duplicates.
"""

PROJECTS_TASK = """### TASK
Return JSON with one field, projects: up to 4 past projects, most recent first (if there are no dates, keep the order in which they are listed). Take them from the ONE-SLIDER first; use the RESUME only if the one-slider describes no projects. Each project has: name, description (ONE sentence, at most 18 words), technologies (at most 4 short tool names) and source ("one_slider" or "resume"). Empty list if no projects are described.
"""

_STR = {"type": "string"}
_STR_LIST = {"type": "array", "items": _STR}

SUMMARY_SCHEMA = {
    "type": "object",
    "properties": {
        "headline": _STR,
        "current_role": _STR,
        "organization_summary": _STR,
        "skills": _STR_LIST,
    },
    "required": ["headline", "current_role", "organization_summary", "skills"],
}

PROJECTS_SCHEMA = {
    "type": "object",
    "properties": {
        "projects": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "name": _STR,
                    "description": _STR,
                    "technologies": _STR_LIST,
                    "source": {"type": "string", "enum": ["one_slider", "resume"]},
                },
                "required": ["name", "description", "technologies", "source"],
            },
        }
    },
    "required": ["projects"],
}

def build_prompt(task: str, resume_text: str, one_slider_text: str) -> str:
    """Documents first (identical for every call), task last."""
    return (
        DOCUMENTS_TEMPLATE.format(
            one_slider=_prepare(one_slider_text, MAX_SLIDER_CHARS) or "Not provided",
            resume=_prepare(resume_text, MAX_RESUME_CHARS) or "Not provided",
        )
        + task
    )


# --------------------------------------------------------------------------
# Certifications: read by code (exact wording, no model time)
# --------------------------------------------------------------------------
_CERT_HEAD = re.compile(
    r"^\W*(?:professional\s+|technical\s+|key\s+)?"
    r"(?:certifications?|certificates?|licen[cs]es?|courses|trainings?)\b",
    re.I,
)
# A heading line consists ONLY of a known heading (so "Project Management Professional" is not one).
_OTHER_HEAD = re.compile(
    r"^\W*(?:education(?:al)?(?:\s+(?:background|qualifications?))?|academics?|"
    r"(?:work\s+|professional\s+|relevant\s+)?experience|employment(?:\s+history)?|"
    r"(?:key\s+|major\s+)?projects?|(?:technical\s+|key\s+|core\s+)?skills?|"
    r"(?:professional\s+)?(?:summary|profile)|objective|achievements?|"
    r"awards?(?:\s*(?:&|and)\s*\w+)?|publications?|languages?|interests?|"
    r"hobbies(?:\s*(?:&|and)\s*interests?)?|references?|declaration|personal(?:\s+details)?)\W*$",
    re.I,
)
_YEAR = re.compile(r"\(?\b((?:19|20)\d{2})\b\)?")
_SPLIT = re.compile(r"\s+[–—-]\s+|\s*\|\s*")
_CERT_WORD = re.compile(r"\b(?:certified|certification|certificate)\b", re.I)
_TRIM = " -–—•*·,;:"


def _cert_section_lines(text: str) -> list[str]:
    """Lines under a 'Certifications' heading, up to the next known heading."""
    lines, grabbing = [], False
    for raw in (text or "").splitlines():
        line = raw.strip()
        if not line:
            continue
        if not grabbing:
            match = _CERT_HEAD.match(line)
            if not match:
                continue
            rest = line[match.end():].strip()
            if len(rest) > 25 and not rest.startswith((":", "-", "–", "—", "|")):
                continue  # a sentence that merely starts with the word
            grabbing = True
            inline = rest.lstrip(":-–—| ").strip()
            if inline and not re.fullmatch(
                r"(?:&|and)?\s*(?:licen[cs]es?|certifications?|certificates?|training)", inline, re.I
            ):  # "Certifications: AWS SAA; Azure DP-100": the list is on the heading line
                lines.extend(p.strip() for p in re.split(r"[;,]", inline) if p.strip())
            continue
        if _OTHER_HEAD.match(line):
            break
        lines.append(line)
        if len(lines) >= 25:
            break
    return lines


def _fallback_cert_lines(text: str) -> list[str]:
    """Used only when no document has a Certifications heading: short lines that say 'certified'."""
    out = []
    for raw in (text or "").splitlines():
        line = raw.strip()
        if _CERT_WORD.search(line) and len(line) <= 100 and len(line.split()) <= 12:
            out.append(line)
    return out[:8]


def _join_wrapped(lines: list[str]) -> list[str]:
    """A line that starts in lower case, or follows a line ending in a comma/dash/'and', continues it."""
    out: list[str] = []
    for line in lines:
        if out and (line[:1].islower() or out[-1].endswith((",", "-", "–", "&", " and", " of"))):
            out[-1] = out[-1] + " " + line
        else:
            out.append(line)
    return out


def _parse_cert(line: str):
    """'Name – Issuer (2023)' -> {name, issuer, year}; None if the line is not a certification."""
    line = " ".join(line.split()).strip(_TRIM)
    year = ""
    found = list(_YEAR.finditer(line))
    if found:
        last = found[-1]
        year = last.group(1)
        line = " ".join((line[: last.start()] + " " + line[last.end():]).split()).strip(_TRIM)
    if line.count("(") != line.count(")"):  # a stray bracket left behind when a year was cut out
        line = line.replace("(", "").replace(")", "")
    parts = [p.strip(_TRIM) for p in _SPLIT.split(line) if p.strip(_TRIM)]
    if not parts:
        return None
    name, issuer = parts[0], ", ".join(parts[1:])
    if len(name) < 3 or len(name) > 150 or len(name.split()) > 20 or not re.search(r"[A-Za-z]", name):
        return None
    return {"name": name, "issuer": issuer if len(issuer) <= 80 else "", "year": year}


def extract_certifications(resume_text: str, one_slider_text: str) -> list[dict]:
    """Certifications exactly as written in the files. No model is used.

    Reads the 'Certifications' section of the one-slider and the resume. If neither document
    has one, falls back to short lines that contain 'certified' or 'certification'.
    """
    sources = [clean_text(one_slider_text), clean_text(resume_text)]
    sectioned = [_cert_section_lines(t) for t in sources]
    if any(sectioned):
        lines = [line for group in sectioned for line in group]
    else:
        lines = [line for t in sources for line in _fallback_cert_lines(t)]

    out, seen = [], set()
    for line in _join_wrapped(lines):
        cert = _parse_cert(line)
        if cert is None:
            continue
        key = " ".join(_words(cert["name"]))
        if key in seen:
            continue
        seen.add(key)
        out.append(cert)
    return out[:MAX_CERTS_READ]

# --------------------------------------------------------------------------
# Talking to the model
# --------------------------------------------------------------------------

def _chat(
    client: ollama.Client, user_prompt: str, schema: dict, model: str, max_output: int, label: str, attempt: int = 1
) -> str:
    options = {
        "temperature": 0,
        "num_ctx": OLLAMA_NUM_CTX,
        "num_predict": min(OLLAMA_MAX_OUTPUT, max_output),
    }
    if OLLAMA_NUM_THREAD > 0:
        options["num_thread"] = OLLAMA_NUM_THREAD
    try:
        response = client.chat(
            model=model,
            messages=[
                {"role": "system", "content": SYSTEM_PROMPT},
                {"role": "user", "content": user_prompt},
            ],
            format=schema,
            keep_alive=OLLAMA_KEEP_ALIVE,
            options=options,
        )
    except ollama.ResponseError as exc:
        if exc.status_code == 404:
            raise SummaryError(f"Model '{model}' is not installed. Run: ollama pull {model}")
        raise SummaryError(f"Ollama error: {exc.error}")
    except (ConnectionError, httpx.ConnectError):
        raise SummaryError("Cannot reach Ollama. Make sure it is running (try: ollama serve).")
    except httpx.TimeoutException:
        raise SummaryError("The model took too long to respond. Try again, or use shorter files.")
    timing.ai_call(label, response, attempt, OLLAMA_NUM_CTX)
    return response["message"]["content"]


def _ask(task: str, schema: dict, max_output: int, label: str, resume_text: str, slider_text: str, model: str) -> dict:
    """One model call -> parsed JSON dict. Retries once if the JSON is unusable."""
    prompt = build_prompt(task, resume_text, slider_text)
    client = ollama.Client(host=OLLAMA_HOST, timeout=OLLAMA_TIMEOUT)
    for attempt in (1, 2):
        content = _chat(client, prompt, schema, model, max_output, label, attempt)
        try:
            data = json.loads(content)
        except json.JSONDecodeError:
            timing.log(f"  {label}: answer was not valid JSON (attempt {attempt})")
            continue
        if isinstance(data, dict):
            return data
    raise SummaryError("The model returned an unusable answer twice. Try again.")


# --------------------------------------------------------------------------
# Cleaning the model's answer
# --------------------------------------------------------------------------
def _as_list(value) -> list[str]:
    if isinstance(value, str):
        value = [value]
    if not isinstance(value, list):
        return []
    return [str(v).strip() for v in value if str(v).strip()]


def _one_line(value, limit: int) -> str:
    return " ".join(str(value or "").split())[:limit]


def _supported(text: str, haystack: str, share: float = 1.0) -> bool:
    """True if at least `share` of the words in `text` occur in the source text
    (a cheap guard against the model inventing a title)."""
    tokens = [t for t in _words(text) if len(t) >= 2]
    if not tokens:
        return False
    found = sum(1 for t in tokens if t in haystack)
    return found / len(tokens) >= share


def _source_text(resume_text: str, slider_text: str) -> str:
    return (clean_text(slider_text) + "\n" + clean_text(resume_text)).lower()


def _clean_projects(value) -> list[dict]:
    out = []
    for item in value if isinstance(value, list) else []:
        if not isinstance(item, dict):
            continue
        name = _one_line(item.get("name"), 120)
        description = _one_line(item.get("description"), 200)
        if not name and not description:
            continue
        source = str(item.get("source") or "").strip().lower()
        out.append(
            {
                "name": name or description[:60],
                "role": _one_line(item.get("role"), 80),
                "description": description,
                "technologies": merge_skills(_as_list(item.get("technologies")))[:5],
                "source": source if source in ("one_slider", "resume") else "resume",
            }
        )
    return out[:4]


# --------------------------------------------------------------------------
# Public functions
# --------------------------------------------------------------------------
def generate_summary_part(resume_text: str, one_slider_text: str, model=None) -> dict:
    """Call 1. Returns {'designation', 'organization_summary', 'skills'}.
    Raises SummaryError with a user-friendly message on any failure."""
    if not (clean_text(resume_text) or clean_text(one_slider_text)):
        raise SummaryError("Nothing to summarize yet. Upload a resume or one-slider first.")
    model = model or models.get_summary_model()
    data = _ask(SUMMARY_TASK, SUMMARY_SCHEMA, SUMMARY_MAX_OUTPUT, "summary", resume_text, one_slider_text, model)

    overview = str(data.get("organization_summary", "")).strip()
    if not overview:
        raise SummaryError("The model returned an empty summary. Try again.")
    haystack = _source_text(resume_text, one_slider_text)
    titles: list[str] = []
    for field_name in ("headline", "current_role"):
        title = _one_line(data.get(field_name), 80)
        if not title or len(title.split()) > 8 or not _supported(title, haystack, 0.5):
            continue  # empty, too long for a title, or not found in the files
        low = title.lower()
        if any(low in t.lower() for t in titles):
            continue  # already covered by a longer title
        titles = [t for t in titles if t.lower() not in low] + [title]
    designation = " - ".join(titles)
    timing.log(f"  designation: {designation!r}")
    return {
        "designation": designation,
        "organization_summary": overview,
        "skills": merge_skills(_as_list(data.get("skills")))[:25],
    }


def generate_projects(resume_text: str, one_slider_text: str, model=None) -> list[dict]:
    """Call 2. Up to 4 projects: {name, role, description, technologies, source}.
    Raises SummaryError on failure."""
    if not (clean_text(resume_text) or clean_text(one_slider_text)):
        return []
    model = model or models.get_summary_model()
    data = _ask(PROJECTS_TASK, PROJECTS_SCHEMA, PROJECTS_MAX_OUTPUT, "projects", resume_text, one_slider_text, model)
    return _clean_projects(data.get("projects"))


def generate_summary(resume_text: str, one_slider_text: str, feedback_list=None, model=None) -> dict:
    """Both calls plus the certification parser, in one result (used by the folder import).

    feedback_list is accepted only so older callers keep working: manager reviews are
    no longer part of the AI summary. If the projects call fails, the key 'projects' is
    left out so stored projects are not overwritten.
    """
    result = generate_summary_part(resume_text, one_slider_text, model)
    try:
        result["projects"] = generate_projects(resume_text, one_slider_text, model)
    except SummaryError as exc:
        timing.log(f"  projects call failed: {exc}")
    with timing.span("certifications (code)"):
        result["certifications"] = extract_certifications(resume_text, one_slider_text)
    result["top_strengths"] = []  # the Strengths section is switched off for now
    return result


def check_ollama() -> tuple[bool, str]:
    """Quick health check for the UI sidebar (uses the currently selected summary model)."""
    installed = models.installed_models()
    if installed is None:
        return False, "Ollama is not reachable. Start it (ollama serve)."
    model = models.get_summary_model()
    if models.is_installed(model, installed):
        return True, f"Ollama ready ({model})"
    return False, f"Model not installed. Run: ollama pull {model}"